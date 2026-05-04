from openpyxl import load_workbook
from inventory.models import AssetMaster

def build_asset_code(data):
    if data.get('num_activo'):
        return str(data['num_activo']).strip()
    return '|'.join(str(data.get(k,'')).strip() for k in ['cuenta','tipo','numero','descripcion','ubicacion'])

def import_assets_from_excel(batch):
    wb = load_workbook(batch.source_file.path, data_only=True)
    specs = [('INC 1 ACT DEPRECIABLES',2),('INC 2 ACT DEPRECIABLES',1)]
    created = 0
    for sheet_name, header_row in specs:
        ws = wb[sheet_name]
        headers = [str(c.value).strip() if c.value else '' for c in ws[header_row]]
        idx = {h:i for i,h in enumerate(headers) if h}
        for r in range(header_row + 1, ws.max_row + 1):
            vals = [ws.cell(r,c).value for c in range(1, ws.max_column +1)]
            if not any(vals):
                continue
            line = ' '.join(str(v or '') for v in vals)
            if 'Total Balanzas' in line:
                continue
            data = {
                'cuenta': vals[idx['Cuenta']] if 'Cuenta' in idx else '',
                'nombre': vals[idx['Nombre']] if 'Nombre' in idx else '',
                'tipo': vals[idx['Tipo']] if 'Tipo' in idx else '',
                'numero': vals[idx['Número']] if 'Número' in idx else '',
                'descripcion': vals[idx['Ch Descripcion']] if 'Ch Descripcion' in idx else '',
                'ubicacion': vals[idx['Ubicación']] if 'Ubicación' in idx else '',
                'num_activo': vals[idx['Num activo']] if 'Num activo' in idx else '',
            }
            asset_code = build_asset_code(data)
            _, was_created = AssetMaster.objects.get_or_create(
                project=batch.project,
                asset_code=asset_code,
                defaults={**data, 'source_sheet': sheet_name, 'source_row': r, 'source_file_name': batch.source_file.name}
            )
            created += int(was_created)
    batch.summary = f'Activos creados: {created}'
    batch.save()
